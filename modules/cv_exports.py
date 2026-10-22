"""Fact-only CV rendering with optional sanitized photos and real document exports."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape


def cv_facts(twin: dict[str, Any]) -> list[dict[str, str]]:
    """Create an allowlist that AI can order but never embellish."""
    profile = twin.get("profile", {})
    facts = []

    def add(section: str, text: str, url: str = "") -> None:
        if text.strip():
            facts.append(
                {
                    "id": f"fact_{len(facts) + 1}",
                    "section": section,
                    "text": text.strip(),
                    "url": url if url.startswith("https://") else "",
                }
            )

    add(
        "Education",
        " | ".join(
            str(profile.get(k) or "")
            for k in ("degree", "institution", "education_level")
            if profile.get(k)
        ),
    )
    if profile.get("graduation_date"):
        add("Education", f"Expected completion: {profile['graduation_date']}")
    for skill in twin.get("skills", []):
        evidence = str(skill.get("evidence", ""))
        label = (
            "Formative assessment evidence"
            if "formative assessment" in evidence
            else "Reported learning"
        )
        add("Skills", f"{skill['name']} - {label}")
    for project in twin.get("projects", []):
        add(
            "Projects",
            (
                " - ".join(
                    str(project.get(k) or "")
                    for k in ("name", "description", "evidence")
                    if project.get(k)
                )
                if isinstance(project, dict)
                else str(project)
            ),
        )
    for course in twin.get("learning_courses", []):
        if course.get("status") == "passed" and not course.get("assessment_offline"):
            attempts = [
                a
                for a in course.get("attempts", [])
                if a.get("passed") and not a.get("practice_only")
            ]
            if attempts:
                add(
                    "Assessed Learning",
                    f"{course['title']} - {course['skill']} | StudentReplica formative assessment {attempts[-1]['score']:g}% | {course.get('passed_at', '')[:10]} | Not an accredited certificate",
                    course.get("url", ""),
                )
    for achievement in twin.get("achievements", []):
        add("Achievements", f"{achievement} (student-reported)")
    for course in twin.get("academic", {}).get("courses", []):
        if course.get("grade") not in (None, ""):
            prefix = (
                f"Semester {course['semester']}: " if course.get("semester") else ""
            )
            add(
                "Relevant Coursework",
                f"{prefix}{course['name']} | Reported grade: {course['grade']}",
            )
    return facts[:80]


def select_facts(facts: list[dict[str, str]], ids: list[str]) -> list[dict[str, str]]:
    """Reject unknown IDs; keep education and passed learning visible by default."""
    by_id = {f["id"]: f for f in facts}
    ordered = list(dict.fromkeys(ids))
    selected = [by_id[i] for i in ordered if i in by_id]
    essential = [
        f
        for f in facts
        if f["section"] in ("Education", "Assessed Learning") and f["id"] not in ordered
    ]
    return essential + selected if selected else facts


def cv_markdown(twin: dict, facts: list[dict]) -> str:
    """Render only literal student facts, with clickable course evidence links."""
    profile = twin.get("profile", {})
    lines = [f"# {profile.get('name') or 'Student'}"]
    location = ", ".join(
        str(profile.get(k) or "") for k in ("city", "country") if profile.get(k)
    )
    if location:
        lines += [location]
    for section in dict.fromkeys(f["section"] for f in facts):
        lines += ["", f"## {section}"]
        for fact in facts:
            if fact["section"] == section:
                lines.append(
                    f"- {fact['text']}"
                    + (f" [Course evidence]({fact['url']})" if fact.get("url") else "")
                )
    return "\n".join(lines)


def sanitize_photo(raw: bytes) -> bytes:
    """Bound image size, remove metadata and normalize the upload before storage."""
    from PIL import Image, ImageOps

    if len(raw) > 3_000_000:
        raise ValueError("Choose a photo under 3 MB.")
    try:
        with Image.open(BytesIO(raw)) as image:
            if image.width * image.height > 20_000_000:
                raise ValueError("Photo dimensions are too large.")
            image = ImageOps.exif_transpose(image).convert("RGB")
            image.thumbnail((600, 600))
            output = BytesIO()
            image.save(output, format="JPEG", quality=85)
            return output.getvalue()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Choose a valid JPG or PNG photo.") from exc


def export_docx(twin: dict, facts: list[dict], photo: bytes | None = None) -> bytes:
    """Build an editable, ATS-readable Word document without invented claims."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    profile = twin.get("profile", {})
    name = profile.get("name") or "Student"

    if photo:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

        header = document.add_table(rows=1, cols=2)
        header.autofit = False
        header.columns[0].width, header.columns[1].width = Inches(5.5), Inches(1)
        title = header.cell(0, 0).paragraphs[0]
        title.style = "Title"
        title.add_run(name)
        portrait = header.cell(0, 1).paragraphs[0]
        portrait.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        portrait.add_run().add_picture(BytesIO(photo), width=Inches(0.85))
        for cell in header.rows[0].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    else:
        title = document.add_paragraph(name, "Title")

    title.runs[0].font.color.rgb = RGBColor(15, 118, 110)
    title.runs[0].font.size = Pt(24)
    title.runs[0].font.bold = True

    sub_parts = [
        str(profile.get(k) or "")
        for k in ("degree", "institution")
        if profile.get(k)
    ]
    if sub_parts:
        sub_p = document.add_paragraph(" • ".join(sub_parts))
        sub_p.runs[0].font.size = Pt(10.5)
        sub_p.runs[0].font.color.rgb = RGBColor(71, 85, 105)
        sub_p.runs[0].font.bold = True

    loc_parts = [
        str(profile.get(k) or "")
        for k in ("city", "country")
        if profile.get(k)
    ]
    if loc_parts:
        loc_p = document.add_paragraph(", ".join(loc_parts))
        loc_p.runs[0].font.size = Pt(9.5)
        loc_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    for heading in dict.fromkeys(f["section"] for f in facts):
        paragraph = document.add_paragraph(heading.upper(), "Heading 1")
        paragraph.runs[0].font.size = Pt(11.5)
        paragraph.runs[0].font.bold = True
        paragraph.runs[0].font.color.rgb = RGBColor(15, 118, 110)
        for fact in (f for f in facts if f["section"] == heading):
            paragraph = document.add_paragraph(fact["text"], "List Bullet")
            if fact.get("url"):
                paragraph.add_run(" ")
                link = OxmlElement("w:hyperlink")
                link.set(
                    qn("r:id"),
                    paragraph.part.relate_to(
                        fact["url"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True
                    ),
                )
                run, text = OxmlElement("w:r"), OxmlElement("w:t")
                text.text = "Course evidence"
                run.append(text)
                link.append(run)
                paragraph._p.append(link)

    # Word templates may inherit a title rule; keep this CV free of theme borders.
    for root in (document.element, document.styles.element):
        for border in list(root.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_pdf(twin: dict, facts: list[dict], photo: bytes | None = None) -> bytes:
    """Render a paginated CV with graphical styling, section hierarchy and clickable links."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, HRFlowable
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font = "Helvetica"
    for path in (
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ):
        if path.exists():
            font = "TwinCV"
            if font not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font, str(path)))
            break

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=48,
        leftMargin=48,
        topMargin=40,
        bottomMargin=40,
        title="Curriculum Vitae",
        author=twin.get("profile", {}).get("name", "Student"),
    )
    body = ParagraphStyle(
        "Body",
        fontName=font,
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155"),
        spaceAfter=5,
        splitLongWords=True,
    )
    heading = ParagraphStyle(
        "SectionHeading",
        parent=body,
        fontSize=11,
        leading=14,
        fontName=f"{font}-Bold" if font == "Helvetica" else font,
        textColor=colors.HexColor("#0f766e"),
        spaceBefore=0,
        spaceAfter=0,
    )
    title = ParagraphStyle(
        "Name",
        parent=body,
        fontSize=24,
        leading=28,
        fontName=f"{font}-Bold" if font == "Helvetica" else font,
        textColor=colors.HexColor("#0f766e"),
        spaceAfter=3,
    )
    subtitle = ParagraphStyle(
        "Subtitle",
        parent=body,
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#475569"),
        spaceAfter=3,
    )
    contact_style = ParagraphStyle(
        "Contact",
        parent=body,
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=0,
    )

    profile = twin.get("profile", {})
    name_p = Paragraph(escape(profile.get("name") or "Student"), title)

    sub_parts = [
        str(profile.get(k) or "")
        for k in ("degree", "institution")
        if profile.get(k)
    ]
    sub_p = Paragraph(escape(" • ".join(sub_parts)), subtitle) if sub_parts else None

    loc_parts = [
        str(profile.get(k) or "")
        for k in ("city", "country")
        if profile.get(k)
    ]
    loc_p = Paragraph(escape(", ".join(loc_parts)), contact_style) if loc_parts else None

    header_cells = [p for p in (name_p, sub_p, loc_p) if p is not None]

    story = []
    if photo:
        image = Image(BytesIO(photo))
        ratio = min(68 / image.imageWidth, 76 / image.imageHeight)
        image.drawWidth, image.drawHeight = (
            image.imageWidth * ratio,
            image.imageHeight * ratio,
        )
        hdr_table = Table([[header_cells, image]], colWidths=[document.width - 80, 80])
        hdr_table.setStyle(
            TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ])
        )
        story.append(hdr_table)
    else:
        for cell in header_cells:
            story.append(cell)

    story.append(
        HRFlowable(
            width="100%",
            thickness=1.5,
            color=colors.HexColor("#0f766e"),
            spaceBefore=6,
            spaceAfter=12,
        )
    )

    for section in dict.fromkeys(f["section"] for f in facts):
        sec_label = Paragraph(f"<b>{escape(section.upper())}</b>", heading)
        sec_table = Table([[sec_label]], colWidths=[document.width])
        sec_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f0fdfa")),
                ("LINELEFT", (0, 0), (0, -1), 3.5, colors.HexColor("#0f766e")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ])
        )
        story.append(sec_table)
        story.append(Spacer(1, 4))
        for fact in (f for f in facts if f["section"] == section):
            text = escape(fact["text"])
            if fact.get("url"):
                text += f' <link href="{escape(fact["url"], {chr(34): "&quot;"})}" color="#0f766e">Course evidence</link>'
            story.append(Paragraph(text, body, bulletText="-"))
        story.append(Spacer(1, 6))

    document.build(story)
    return output.getvalue()
